"""
Loads text out of uploaded medical documents.

Supports:
  - Digital PDFs (text extracted directly via pypdf)
    - Scanned / photographed PDFs (OCR fallback via PDFium + pytesseract)
  - Word documents (.docx)
  - Plain images of prescriptions (.png/.jpg/.jpeg) via OCR

If a PDF page yields almost no extractable text (a strong signal it's a
scanned image, e.g. a photographed handwritten prescription), we
automatically fall back to OCR for that page.
"""
import os
from pathlib import Path
from pypdf import PdfReader
from docx import Document as DocxDocument
from PIL import Image
import pytesseract
import pypdfium2 as pdfium

TESSERACT_CMD = Path(os.getenv(
    "TESSERACT_CMD",
    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
))
if TESSERACT_CMD.is_file():
    pytesseract.pytesseract.tesseract_cmd = str(TESSERACT_CMD)

MIN_CHARS_PER_PAGE = 30  # below this, assume the page is a scanned image


def _ocr_image(image: Image.Image) -> str:
    try:
        return pytesseract.image_to_string(image)
    except Exception as e:
        raise RuntimeError(f"OCR failed: {e}") from e


def _load_pdf(path: Path) -> str:
    text_parts = []
    reader = PdfReader(str(path))
    pages_needing_ocr = []

    for i, page in enumerate(reader.pages):
        page_text = (page.extract_text() or "").strip()
        if len(page_text) >= MIN_CHARS_PER_PAGE:
            text_parts.append(page_text)
        else:
            text_parts.append(None)  # placeholder, fill in via OCR below
            pages_needing_ocr.append(i)

    if pages_needing_ocr:
        try:
            pdf_document = pdfium.PdfDocument(str(path))
            try:
                for i in pages_needing_ocr:
                    pdf_page = pdf_document[i]
                    try:
                        image = pdf_page.render(scale=2).to_pil()
                        text_parts[i] = _ocr_image(image)
                    finally:
                        pdf_page.close()
            finally:
                pdf_document.close()
        except Exception as e:
            raise RuntimeError(f"Could not OCR scanned PDF pages: {e}") from e

    return "\n\n".join(t for t in text_parts if t)


def _load_docx(path: Path) -> str:
    doc = DocxDocument(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also pull text out of any tables (common in lab reports)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip(" |"):
                parts.append(row_text)

    return "\n".join(parts)


def _load_image(path: Path) -> str:
    image = Image.open(path)
    return _ocr_image(image)


def load_document(path: Path) -> str:
    """Dispatch by file extension and return extracted plain text."""
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return _load_pdf(path)
    elif suffix in (".docx", ".doc"):
        return _load_docx(path)
    elif suffix in (".png", ".jpg", ".jpeg"):
        return _load_image(path)
    else:
        raise ValueError(f"Unsupported file type: {suffix}")
